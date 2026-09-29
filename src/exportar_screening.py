"""Exportação da tabela de screening para Excel, formatada e ordenável."""

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from config.config import PROCESSED_DIR
from src.exportar import _caminho_longo, COR_TOTAL

COLUNAS_FINAIS = {
    "DENOM_SOCIAL": "Empresa",
    "SETOR_ATIV": "Setor",
    "ticker_on": "Ticker ON",
    "ticker_pn": "Ticker PN",
    "preco": "Preço (R$)",
    "market_cap": "Market Cap (R$)",
    "p_l": "P/L",
    "p_vp": "P/VP",
    "dividend_yield_pct": "Dividend Yield (%)",
    "roe_pct": "ROE (%)",
    "margem_bruta_pct": "Margem Bruta (%)",
    "margem_liquida_pct": "Margem Líquida (%)",
    "crescimento_receita_pct": "Crescimento Receita (%)",
    "liquidez_corrente": "Liquidez Corrente",
    "volume_medio_30d": "Volume Médio 30d",
}

FORMATOS = {
    "Preço (R$)": "#,##0.00",
    "Market Cap (R$)": "#,##0",
    "P/L": "0.00",
    "P/VP": "0.00",
    "Dividend Yield (%)": "0.00",
    "ROE (%)": "0.00",
    "Margem Bruta (%)": "0.00",
    "Margem Líquida (%)": "0.00",
    "Crescimento Receita (%)": "0.00",
    "Liquidez Corrente": "0.00",
    "Volume Médio 30d": "#,##0",
}


def exportar_screening(tabela: pd.DataFrame, ano_fiscal: int = None) -> Path:
    """
    Exporta a tabela de screening (saída de montar_screening()) pra Excel,
    uma linha por empresa, pronta pra ordenar/filtrar no próprio Excel.
    """
    ano_fiscal = ano_fiscal or tabela.attrs.get("ano_fiscal")
    colunas_presentes = [c for c in COLUNAS_FINAIS if c in tabela.columns]
    saida = tabela[colunas_presentes].rename(columns=COLUNAS_FINAIS)
    saida = saida.sort_values("Empresa").reset_index(drop=True)

    caminho = PROCESSED_DIR / f"Screening_Brasil_{ano_fiscal}_{date.today():%Y%m%d}.xlsx"

    with pd.ExcelWriter(_caminho_longo(caminho), engine="openpyxl") as writer:
        saida.to_excel(writer, sheet_name="Screening", index=False)

    _formatar(caminho)
    print(f"\n📁 Screening salvo: {caminho} ({len(saida)} empresas)")
    return caminho


def _formatar(caminho: Path):
    wb = load_workbook(_caminho_longo(caminho))
    ws = wb["Screening"]

    fill_header = PatternFill("solid", fgColor=COR_TOTAL)
    fonte_header = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill_header
        cell.font = fonte_header
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    cabecalhos = [c.value for c in ws[1]]
    for col_idx, nome in enumerate(cabecalhos, start=1):
        if nome in FORMATOS:
            letra = get_column_letter(col_idx)
            for cell in ws[letra][1:]:
                cell.number_format = FORMATOS[nome]

    larguras = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            letra = get_column_letter(cell.column)
            larguras[letra] = max(larguras.get(letra, 10), len(str(cell.value)))
    for letra, tamanho in larguras.items():
        ws.column_dimensions[letra].width = min(tamanho + 2, 40)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(_caminho_longo(caminho))
