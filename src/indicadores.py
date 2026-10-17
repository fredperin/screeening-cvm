"""
indicadores.py — Aba "Indicadores (LTM)" do documento padrão.

Todas as linhas são FÓRMULAS do Excel ligadas à aba "Demonstrações" (nada é
colado como número), então o usuário audita clicando na célula. Fluxos (DRE e
DFC) são sempre somados nos últimos 4 trimestres (LTM = trailing 12 meses);
saldos (Balanço) são os do fim do período.

Contas até o nível 3 do plano da CVM são padronizadas entre empresas; D&A e
Capex vêm de contas de nível 4 (que variam por empresa) e por isso são
identificados pela DESCRIÇÃO — aproximação, listada na seção "Notas".
"""

import re
from datetime import datetime

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COR_TITULO = "1A531A"
COR_SECAO = "EAC037"
FMT_NUM = "#,##0;(#,##0)"
FMT_PCT = "0.0%;(0.0%)"
FMT_X = '0.00"x";(0.00"x")'

ABA_DEMONSTRACOES = "Demonstrações"
ABA_INDICADORES = "Indicadores (LTM)"

PADRAO_DA = re.compile(r"deprecia|amortiza|exaust", re.I)
EXCLUI_DA = re.compile(r"juros|encargos|imposto|tribut|ágio|agio|custo de transa|variaç", re.I)
PADRAO_CAPEX = re.compile(r"imobilizado|intang[ií]vel", re.I)
EXCLUI_CAPEX = re.compile(r"venda|alienaç|baixa|receb|produto d|resgate|dividendo", re.I)


class _Bloco:
    """Posição de um bloco (DRE/Balanço/DFC) dentro da aba Demonstrações."""

    def __init__(self, df, linha_dados):
        self.codigos = df.attrs.get("codigos", [])
        self.descricoes = list(df.index)
        self.rotulos = list(df.columns)
        self.linha_dados = linha_dados

    def linha(self, cd):
        return self.linha_dados + self.codigos.index(cd) if cd in self.codigos else None

    def linhas_por_descricao(self, prefixo, padrao, exclui):
        return [
            self.linha_dados + i
            for i, (cd, desc) in enumerate(zip(self.codigos, self.descricoes))
            if cd.startswith(prefixo + ".") and cd.count(".") == prefixo.count(".") + 1
            and padrao.search(str(desc)) and not exclui.search(str(desc))
        ]


def _data(rotulo):
    return datetime.strptime(rotulo, "%d/%m/%y")


def _janela_trimestres_consecutivos(rotulos, j):
    """True se os 4 rótulos terminando em j são trimestres seguidos (sem buraco no meio)."""
    if j < 3:
        return False
    datas = [_data(r) for r in rotulos[j - 3: j + 1]]
    return all(70 <= (b - a).days <= 105 for a, b in zip(datas, datas[1:]))


def adicionar_aba_indicadores(wb, blocos_posicionados, blocos, ano_exibicao=None):
    """
    Args:
        wb:                   workbook openpyxl já com a aba "Demonstrações"
        blocos_posicionados:  lista (linha_titulo, titulo, linha_dados, estilos) da exportação
        blocos:               lista (titulo, DataFrame) — saída de montar_documento_padrao()
        ano_exibicao:         primeiro ano mostrado na aba; períodos anteriores existem só pra
                              alimentar o LTM dos primeiros trimestres e não viram coluna
    """
    linha_dados = {titulo: linha for _, titulo, linha, _ in blocos_posicionados}
    por_titulo = {t: df for t, df in blocos if df is not None and not df.empty and "codigos" in df.attrs}

    dre = _Bloco(por_titulo["DRE"], linha_dados["DRE"]) if "DRE" in por_titulo else None
    bp = _Bloco(por_titulo["Balanço Patrimonial"], linha_dados["Balanço Patrimonial"]) if "Balanço Patrimonial" in por_titulo else None
    dfc = _Bloco(por_titulo["DFC (Método Indireto)"], linha_dados["DFC (Método Indireto)"]) if "DFC (Método Indireto)" in por_titulo else None

    if not (dre and bp and dfc):
        print("  ⚠️ Aba de indicadores não gerada: falta DRE, Balanço ou DFC")
        return
    if not (dre.rotulos == bp.rotulos == dfc.rotulos):
        print("  ⚠️ Aba de indicadores não gerada: os períodos de DRE, Balanço e DFC não coincidem")
        return

    rotulos = dre.rotulos
    exibidos = [j for j, r in enumerate(rotulos) if not ano_exibicao or _data(r).year >= ano_exibicao]
    posicao = {j: k for k, j in enumerate(exibidos)}  # j (aba Demonstrações) -> coluna na aba de indicadores
    n = len(exibidos)
    ws = wb.create_sheet(ABA_INDICADORES)
    aba = f"'{ABA_DEMONSTRACOES}'!"

    def col(j):
        return get_column_letter(j + 2)  # coluna na aba Demonstrações (B = primeiro período)

    def col_ind(j):
        return get_column_letter(posicao[j] + 2)  # coluna na aba de indicadores

    def pontual(bloco, cd, j):
        linha = bloco.linha(cd)
        return f"{aba}{col(j)}{linha}" if linha else None

    def ltm(bloco, cd, j):
        linha = bloco.linha(cd)
        return f"SUM({aba}{col(j - 3)}{linha}:{col(j)}{linha})" if linha else None

    def ltm_linhas(bloco, linhas, j):
        if not linhas:
            return None
        return "SUM(" + ",".join(f"{aba}{col(j - 3)}{r}:{col(j)}{r}" for r in linhas) + ")"

    linhas_da = dfc.linhas_por_descricao("6.01.01", PADRAO_DA, EXCLUI_DA)
    linhas_capex = dfc.linhas_por_descricao("6.02", PADRAO_CAPEX, EXCLUI_CAPEX)

    # --- cabeçalho -----------------------------------------------------
    ws["A1"] = "Indicadores — sempre trailing 12 meses (LTM)"
    ws["A2"] = "Alíquota de IR/CSLL para o NOPAT (editável)"
    ws["B2"] = 0.34
    ws["B2"].number_format = "0%"
    ws["B2"].font = Font(bold=True, color="0000FF")
    ALIQ = "$B$2"

    cabecalho = 4
    ws.cell(cabecalho, 1, "Período (fim do trimestre)")
    for j in exibidos:
        ws.cell(cabecalho, posicao[j] + 2, rotulos[j])

    linhas_saida = []  # (rótulo, tipo, formato, fn(j) -> fórmula | None)
    ref = {}           # rótulo -> número da linha na aba de indicadores

    def secao(titulo):
        linhas_saida.append((titulo, "secao", None, None))

    def linha(rotulo, formato, fn):
        linhas_saida.append((rotulo, "dado", formato, fn))

    def c(rotulo, j):
        return f"{col_ind(j)}{ref[rotulo]}"

    def div(a, b):
        return f'IFERROR({a}/{b},"n/d")'

    # --- seção 1: base LTM --------------------------------------------
    secao("BASE — acumulado dos últimos 12 meses (R$, mesma unidade da aba Demonstrações)")
    linha("Receita Líquida", FMT_NUM, lambda j: ltm(dre, "3.01", j))
    linha("Lucro Bruto", FMT_NUM, lambda j: ltm(dre, "3.03", j))
    linha("EBIT (Resultado antes do resultado financeiro e tributos)", FMT_NUM, lambda j: ltm(dre, "3.05", j))
    linha("Depreciação e Amortização (aprox., ver Notas)", FMT_NUM, lambda j: ltm_linhas(dfc, linhas_da, j) or "0")
    linha("EBITDA", FMT_NUM, lambda j: f"{c('EBIT (Resultado antes do resultado financeiro e tributos)', j)}+{c('Depreciação e Amortização (aprox., ver Notas)', j)}")
    linha("Resultado Financeiro", FMT_NUM, lambda j: ltm(dre, "3.06", j))
    linha("Lucro Antes dos Tributos", FMT_NUM, lambda j: ltm(dre, "3.07", j))
    linha("Lucro Líquido (consolidado)", FMT_NUM, lambda j: ltm(dre, "3.11", j))
    linha("Fluxo de Caixa Operacional (FCO)", FMT_NUM, lambda j: ltm(dfc, "6.01", j))
    linha("Capex (aprox., ver Notas)", FMT_NUM, lambda j: (f"-{ltm_linhas(dfc, linhas_capex, j)}" if linhas_capex else "0"))
    linha("Fluxo de Caixa Livre (FCO − Capex)", FMT_NUM, lambda j: f"{c('Fluxo de Caixa Operacional (FCO)', j)}-{c('Capex (aprox., ver Notas)', j)}")

    # --- seção 2: posição ----------------------------------------------
    secao("POSIÇÃO — saldo no fim de cada período")
    linha("Caixa e Equivalentes", FMT_NUM, lambda j: pontual(bp, "1.01.01", j) or "0")
    linha("Aplicações Financeiras", FMT_NUM, lambda j: pontual(bp, "1.01.02", j) or "0")
    linha("Dívida Bruta (empréstimos CP + LP)", FMT_NUM,
          lambda j: "+".join(x for x in (pontual(bp, "2.01.04", j), pontual(bp, "2.02.01", j)) if x) or "0")
    linha("Dívida Líquida", FMT_NUM,
          lambda j: f"{c('Dívida Bruta (empréstimos CP + LP)', j)}-{c('Caixa e Equivalentes', j)}-{c('Aplicações Financeiras', j)}")
    linha("Patrimônio Líquido (incl. não controladores)", FMT_NUM, lambda j: pontual(bp, "2.03", j))
    linha("Capital Investido (Dívida Líquida + PL)", FMT_NUM,
          lambda j: f"{c('Dívida Líquida', j)}+{c('Patrimônio Líquido (incl. não controladores)', j)}")

    # --- seção 3: indicadores -------------------------------------------
    secao("INDICADORES")
    linha("Crescimento da Receita LTM (vs. mesmo período 1 ano antes)", FMT_PCT,
          lambda j: (f'IFERROR({c("Receita Líquida", j)}/{ltm(dre, "3.01", j - 4)}-1,"n/d")'
                     if _janela_trimestres_consecutivos(rotulos, j - 4) else None))
    linha("Margem Bruta", FMT_PCT, lambda j: div(c("Lucro Bruto", j), c("Receita Líquida", j)))
    linha("Margem EBITDA", FMT_PCT, lambda j: div(c("EBITDA", j), c("Receita Líquida", j)))
    linha("Margem EBIT", FMT_PCT, lambda j: div(c("EBIT (Resultado antes do resultado financeiro e tributos)", j), c("Receita Líquida", j)))
    linha("Margem Líquida", FMT_PCT, lambda j: div(c("Lucro Líquido (consolidado)", j), c("Receita Líquida", j)))
    linha("Margem FCF (FCF / Receita)", FMT_PCT, lambda j: div(c("Fluxo de Caixa Livre (FCO − Capex)", j), c("Receita Líquida", j)))
    linha("Dívida Líquida / EBITDA", FMT_X, lambda j: div(c("Dívida Líquida", j), c("EBITDA", j)))
    linha("Dívida Líquida / Patrimônio Líquido", FMT_X, lambda j: div(c("Dívida Líquida", j), c("Patrimônio Líquido (incl. não controladores)", j)))
    linha("Liquidez Corrente", FMT_X, lambda j: div(pontual(bp, "1.01", j), pontual(bp, "2.01", j)))
    linha("ROE (Lucro Líquido LTM / PL final)", FMT_PCT, lambda j: div(c("Lucro Líquido (consolidado)", j), c("Patrimônio Líquido (incl. não controladores)", j)))
    linha("ROIC (NOPAT LTM / Capital Investido final)", FMT_PCT,
          lambda j: div(f"{c('EBIT (Resultado antes do resultado financeiro e tributos)', j)}*(1-{ALIQ})", c("Capital Investido (Dívida Líquida + PL)", j)))
    linha("Conversão de Caixa (FCO / EBITDA)", FMT_PCT, lambda j: div(c("Fluxo de Caixa Operacional (FCO)", j), c("EBITDA", j)))
    linha("Conversão de Lucro (FCF / Lucro Líquido)", FMT_PCT, lambda j: div(c("Fluxo de Caixa Livre (FCO − Capex)", j), c("Lucro Líquido (consolidado)", j)))

    # --- escrita --------------------------------------------------------
    fill_titulo = PatternFill("solid", fgColor=COR_TITULO)
    fill_secao = PatternFill("solid", fgColor=COR_SECAO)
    linha_atual = cabecalho + 1
    for rotulo, tipo, formato, fn in linhas_saida:
        if tipo == "secao":
            for k in range(1, n + 2):
                ws.cell(linha_atual, k).fill = fill_secao
            ws.cell(linha_atual, 1, rotulo).font = Font(bold=True)
        else:
            ref[rotulo] = linha_atual
            ws.cell(linha_atual, 1, rotulo)
        linha_atual += 1

    linha_atual = cabecalho + 1
    for rotulo, tipo, formato, fn in linhas_saida:
        if tipo == "dado":
            for j in exibidos:
                if not _janela_trimestres_consecutivos(rotulos, j):
                    continue
                formula = fn(j)
                if formula:
                    cell = ws.cell(linha_atual, posicao[j] + 2, f"={formula}")
                    cell.number_format = formato
        linha_atual += 1

    # --- notas ----------------------------------------------------------
    linha_atual += 1
    ws.cell(linha_atual, 1, "NOTAS E CRITÉRIOS").font = Font(bold=True)
    for k in range(1, n + 2):
        ws.cell(linha_atual, k).fill = fill_secao
    notas = [
        "LTM = soma dos últimos 4 trimestres seguidos da DRE/DFC; fica em branco se a empresa não publicou algum desses trimestres.",
        "Saldos (caixa, dívida, PL) são os do fim do trimestre. 'n/d' = divisão por zero.",
        "Dívida Bruta = Empréstimos e Financiamentos do Passivo Circulante (2.01.04) + Não Circulante (2.02.01). Não inclui arrendamentos nem debêntures lançadas em outras contas.",
        "ROIC = EBIT × (1 − alíquota da célula B2) ÷ (Dívida Líquida + PL) do fim do período — não usa capital médio.",
        "EBITDA = EBIT + D&A. D&A e Capex NÃO são contas padronizadas da CVM: são somadas por descrição na DFC e podem divergir do número divulgado pela empresa.",
        "D&A (DFC 6.01.01): " + ("; ".join(dfc.descricoes[r - dfc.linha_dados] for r in linhas_da) or "nenhuma linha encontrada — D&A = 0 e o EBITDA fica igual ao EBIT"),
        "Capex (DFC 6.02, sinal invertido): " + ("; ".join(dfc.descricoes[r - dfc.linha_dados] for r in linhas_capex) or "nenhuma linha encontrada — Capex = 0"),
    ]
    for nota in notas:
        linha_atual += 1
        ws.cell(linha_atual, 1, nota).alignment = Alignment(wrap_text=False)

    # --- formatação -----------------------------------------------------
    ws["A1"].font = Font(bold=True, size=13)
    for k in range(1, n + 2):
        cell = ws.cell(cabecalho, k)
        cell.fill = fill_titulo
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center" if k > 1 else "left")
    ws.column_dimensions["A"].width = 62
    for j in exibidos:
        ws.column_dimensions[col_ind(j)].width = 14
    ws.freeze_panes = ws.cell(cabecalho + 1, 2)
    print(f"  ✅ Aba '{ABA_INDICADORES}' criada ({len(linhas_saida)} linhas, {n} períodos)")


# --- valores em cache ------------------------------------------------------
# O openpyxl grava a fórmula sem o resultado. Excel em Modo de Exibição Protegido
# (todo arquivo baixado), visualizadores web e celular não recalculam, então as
# células apareceriam vazias. Aqui o resultado é calculado em Python (a gramática
# das fórmulas desta aba é pequena: SUM, + - * /, IFERROR) e gravado junto.

_REF_DEM = re.compile(r"'Demonstrações'!\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
_REF_LOCAL = re.compile(r"(?<![A-Za-z!'])\$?([A-Z]{1,3})\$?(\d+)(?![\d(])")


def _valores_em_cache(wb):
    from openpyxl.utils import column_index_from_string

    dem, ind = wb[ABA_DEMONSTRACOES], wb[ABA_INDICADORES]
    memo = {}

    def num(v):
        return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0

    def celula_dem(col, lin):
        return num(dem[f"{col}{lin}"].value)

    def intervalo_dem(c1, l1, c2, l2):
        return [num(dem.cell(r, c).value)
                for r in range(int(l1), int(l2) + 1)
                for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1)]

    def celula_ind(ref):
        if ref not in memo:
            valor = ind[ref].value
            memo[ref] = avaliar(valor[1:]) if isinstance(valor, str) and valor.startswith("=") else num(valor)
        return memo[ref]

    def soma(*args):
        return sum(sum(a) if isinstance(a, list) else a for a in args)

    def iferror(fn, alternativa):
        try:
            return fn()
        except (ZeroDivisionError, TypeError):
            return alternativa

    def avaliar(formula):
        def dem_ref(m):
            c1, l1, c2, l2 = m.groups()
            if c2:
                return f"_I('{c1}',{l1},'{c2}',{l2})"
            return f"_D('{c1}',{l1})"
        expr = _REF_DEM.sub(dem_ref, formula)
        partes = re.split(r"(_[ID]\([^)]*\))", expr)  # não mexe nas referências já convertidas
        expr = "".join(p if p.startswith(("_I(", "_D(")) else _REF_LOCAL.sub(lambda m: f"_L('{m.group(1)}{m.group(2)}')", p) for p in partes)
        expr = expr.replace("SUM(", "_S(").replace("IFERROR(", "_E(lambda:")
        return eval(expr, {"__builtins__": {}, "_D": celula_dem, "_I": intervalo_dem, "_L": celula_ind, "_S": soma, "_E": iferror})

    resultados = {}
    for linha in ind.iter_rows():
        for cell in linha:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                resultados[cell.coordinate] = celula_ind(cell.coordinate)
    return resultados


def gravar_valores_em_cache(caminho, wb):
    """Reabre o .xlsx salvo e injeta o resultado calculado de cada fórmula da aba de indicadores."""
    import zipfile
    from xml.sax.saxutils import escape

    if ABA_INDICADORES not in wb.sheetnames:
        return  # aba não foi gerada (ex: empresa sem DFC indireta)
    valores = _valores_em_cache(wb)
    indice = wb.sheetnames.index(ABA_INDICADORES) + 1
    nome_xml = f"xl/worksheets/sheet{indice}.xml"

    def injetar(m):
        ref, atributos, formula = m.group(1), m.group(2) or "", m.group(3)
        valor = valores.get(ref)
        if valor is None:
            return m.group(0)
        if isinstance(valor, str):
            return f'<c r="{ref}"{atributos} t="str"><f>{formula}</f><v>{escape(valor)}</v></c>'
        return f'<c r="{ref}"{atributos}><f>{formula}</f><v>{valor!r}</v></c>'

    with zipfile.ZipFile(caminho) as origem:
        arquivos = {n: origem.read(n) for n in origem.namelist()}
    xml = arquivos[nome_xml].decode("utf-8")
    xml = re.sub(r'<c r="([A-Z]+\d+)"( s="\d+")?><f>(.*?)</f><v></v></c>', injetar, xml)
    arquivos[nome_xml] = xml.encode("utf-8")
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as destino:
        for nome, dados in arquivos.items():
            destino.writestr(nome, dados)
