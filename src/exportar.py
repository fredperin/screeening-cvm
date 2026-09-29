"""Exportação das demonstrações CVM para uma planilha Excel formatada."""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from config.config import PROCESSED_DIR

COR_HEADER = "1F4E79"


def _caminho_longo(caminho: Path) -> str:
    """
    No Windows, caminhos com mais de ~260 caracteres (comum aqui, já que o
    workspace fica em uma pasta bem aninhada) quebram funções de I/O padrão.
    O prefixo \\\\?\\ pede ao Windows para ignorar esse limite.
    """
    resolvido = str(caminho.resolve())
    if sys.platform.startswith("win") and not resolvido.startswith("\\\\?\\"):
        return "\\\\?\\" + resolvido
    return resolvido


def exportar_itr(demonstracoes: dict, nome_empresa: str, ano_inicio: int, ano_fim: int = None) -> Path:
    """
    Exporta as demonstrações do ITR (já pivotadas) para um único Excel,
    uma aba por demonstração.

    Args:
        demonstracoes: dict {"Balanço Patrimonial Ativo": df, ...} — saída de buscar_itr()
        nome_empresa:  usado no nome do arquivo
        ano_inicio:    primeiro ano buscado (usado no nome do arquivo)
        ano_fim:       último ano buscado; se None ou igual a ano_inicio, mostra só um ano
    Returns:
        Path do arquivo gerado
    """
    slug = "".join(c if c.isalnum() else "_" for c in nome_empresa).strip("_")
    slug = slug[:25].rstrip("_")  # caminhos muito longos estouram o limite do Windows
    sufixo_ano = str(ano_inicio) if not ano_fim or ano_fim == ano_inicio else f"{ano_inicio}_{ano_fim}"
    caminho = PROCESSED_DIR / f"{slug}_ITR_{sufixo_ano}.xlsx"

    with pd.ExcelWriter(_caminho_longo(caminho), engine="openpyxl") as writer:
        for titulo, df in demonstracoes.items():
            if df is None or df.empty:
                continue
            aba = titulo[:31]
            df.to_excel(writer, sheet_name=aba)
            print(f"  ✅ Aba '{aba}' exportada ({len(df)} linhas)")

    _formatar(caminho)
    print(f"\n📁 Arquivo salvo: {caminho}")
    return caminho


# Paleta do padrão visual do projeto (definida a partir do exemplo formatado
# manualmente pelo usuário — ver src/padrao_visual.py para a classificação
# total/subtotal/detalhe de cada conta).
COR_TOTAL = "1A531A"      # verde — resultados/saldos-chave (título das seções também)
COR_SUBTOTAL = "BFBFBF"   # cinza — agregações estruturais (Circulante, etc.)
FORMATO_NUMERO = "#,##0;(#,##0)"

# Subconjunto do cadastro geral da CVM (46 campos ao todo) relevante pra
# investidor — fora endereço/telefone/fax e os dados de contato do
# responsável, que não agregam à análise.
ROTULOS_CADASTRO = {
    "DENOM_SOCIAL": "Razão Social",
    "DENOM_COMERC": "Nome Comercial",
    "Nome_Empresarial_Anterior": "Nome Empresarial Anterior",  # FCA
    "CNPJ_CIA": "CNPJ",
    "CD_CVM": "Código CVM",
    "Codigo_Negociacao": "Ticker(s) de Negociação",  # FCA
    "Segmento": "Segmento de Listagem (B3)",  # FCA
    "SETOR_ATIV": "Setor de Atividade (CVM)",
    "Descricao_Atividade": "Descrição do Negócio",  # FCA
    "SIT": "Situação",
    "DT_INI_SIT": "Situação desde",
    "TP_MERC": "Tipo de Mercado",
    "CATEG_REG": "Categoria de Registro",
    "DT_INI_CATEG": "Categoria desde",
    "CONTROLE_ACIONARIO": "Controle Acionário",
    "DT_REG": "Data de Registro na CVM",
    "DT_CONST": "Data de Constituição",
    "MUN": "Município (sede)",
    "UF": "UF (sede)",
    "PAIS": "País",
    "CNPJ_AUDITOR": "CNPJ do Auditor",
    "AUDITOR": "Auditor Independente",
    "RESP": "Diretor de Relações com Investidores",
    "EMAIL": "E-mail (RI)",
    "Pagina_Web": "Site",  # FCA
}


def exportar_documento_padrao(
    blocos: list, nome_empresa: str, ano_inicio: int, ano_fim: int = None, dados_cadastrais: dict = None,
) -> Path:
    """
    Exporta o documento padrão do projeto: DRE, Balanço Patrimonial e DFC
    empilhados em UMA ÚNICA ABA, nessa ordem, cada um com seu próprio
    cabeçalho de período (as datas/intervalos não são os mesmos entre as
    três demonstrações, então cada bloco mantém suas próprias colunas).

    Aplica o padrão visual do projeto: linhas "total" e "subtotal" (definidas
    em df.attrs["estilos"], vindo de montar_documento_padrao) recebem
    destaque de cor; números em milhar com negativos entre parênteses.

    Args:
        blocos:           lista de tuplas (titulo, DataFrame) — saída de montar_documento_padrao()
        nome_empresa:     usado no nome do arquivo
        ano_inicio:       primeiro ano buscado (nome do arquivo)
        ano_fim:          último ano buscado (nome do arquivo)
        dados_cadastrais: opcional — saída de buscar_dados_cadastrais(); se informado,
                          vira uma aba "Cadastro" separada (setor, situação, tipo de
                          mercado, controle acionário, auditor etc. — não é dado
                          financeiro, por isso fica fora da aba "Demonstrações")
    Returns:
        Path do arquivo gerado
    """
    slug = "".join(c if c.isalnum() else "_" for c in nome_empresa).strip("_")[:25].rstrip("_")
    sufixo_ano = str(ano_inicio) if not ano_fim or ano_fim == ano_inicio else f"{ano_inicio}_{ano_fim}"
    caminho = PROCESSED_DIR / f"{slug}_Padrao_{sufixo_ano}.xlsx"
    aba = "Demonstrações"

    # (linha_do_titulo_1_indexed, titulo, linha_inicio_dados_1_indexed, estilos_por_linha)
    blocos_posicionados = []
    linha = 0  # 0-indexed, no referencial do pandas (startrow)

    with pd.ExcelWriter(_caminho_longo(caminho), engine="openpyxl") as writer:
        if dados_cadastrais:
            linhas_cadastro = [
                (rotulo, dados_cadastrais.get(campo))
                for campo, rotulo in ROTULOS_CADASTRO.items()
                if pd.notna(dados_cadastrais.get(campo)) and str(dados_cadastrais.get(campo)).strip()
            ]
            if linhas_cadastro:
                df_cadastro = pd.DataFrame(linhas_cadastro, columns=["Campo", "Valor"])
                df_cadastro.to_excel(writer, sheet_name="Cadastro", index=False)
                print(f"  ✅ Aba 'Cadastro' exportada ({len(df_cadastro)} campos)")

        for titulo, df in blocos:
            if df is None or df.empty:
                print(f"  ⚠️ Bloco '{titulo}' vazio, pulando")
                continue

            estilos = df.attrs.get("estilos", [])

            # O título ocupa a MESMA linha do cabeçalho de período (o nome do
            # índice do DataFrame vira o texto do título na coluna A). Os
            # dados começam logo na linha seguinte, sem faixa divisória.
            df = df.copy()
            df.index.name = titulo
            df.to_excel(writer, sheet_name=aba, startrow=linha, index=True)

            linha_titulo = linha + 1   # 1-indexed: título + cabeçalho de período
            linha_dados = linha_titulo + 1
            blocos_posicionados.append((linha_titulo, titulo, linha_dados, estilos))
            print(f"  ✅ Bloco '{titulo}' exportado ({len(df)} linhas)")

            # +1 (título/cabeçalho) + len(df) (dados) + 2 (linhas em branco de separação)
            linha += 1 + len(df) + 2

    _formatar_documento_padrao(caminho, aba, blocos_posicionados)
    print(f"\n📁 Arquivo salvo: {caminho}")
    return caminho


def _formatar_aba_cadastro(ws):
    """Cabeçalho verde/negrito + colunas ajustadas, igual ao padrão das outras abas."""
    fill_total = PatternFill("solid", fgColor=COR_TOTAL)
    fonte_titulo = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = fill_total
        cell.font = fonte_titulo

    larguras = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            letra = get_column_letter(cell.column)
            larguras[letra] = max(larguras.get(letra, 10), len(str(cell.value)))
    for letra, tamanho in larguras.items():
        ws.column_dimensions[letra].width = min(tamanho + 3, 60)

    ws.freeze_panes = "A2"


def _formatar_documento_padrao(caminho: Path, aba: str, blocos_posicionados: list):
    wb = load_workbook(_caminho_longo(caminho))

    if "Cadastro" in wb.sheetnames:
        _formatar_aba_cadastro(wb["Cadastro"])

    ws = wb[aba]

    fill_total = PatternFill("solid", fgColor=COR_TOTAL)
    fill_subtotal = PatternFill("solid", fgColor=COR_SUBTOTAL)
    fonte_titulo = Font(color="FFFFFF", bold=True, size=12)
    fonte_total = Font(color="FFFFFF", bold=True)
    fonte_subtotal = Font(bold=True)

    for linha_titulo, titulo, linha_dados, estilos in blocos_posicionados:
        # Título + cabeçalho de período (mesma linha)
        for cell in ws[linha_titulo]:
            if isinstance(cell, MergedCell):
                continue
            cell.fill = fill_total
            cell.font = fonte_titulo
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        # Linhas de dados: cor por classificação + formato numérico
        for i, estilo in enumerate(estilos):
            linha_excel = linha_dados + i
            linha_celulas = ws[linha_excel]

            if estilo == "total":
                fill, fonte = fill_total, fonte_total
            elif estilo == "subtotal":
                fill, fonte = fill_subtotal, fonte_subtotal
            else:
                fill, fonte = None, None

            for cell in linha_celulas:
                if isinstance(cell, MergedCell):
                    continue
                if cell.column > 1:
                    cell.number_format = FORMATO_NUMERO
                if fill is not None:
                    cell.fill = fill
                if fonte is not None:
                    cell.font = fonte

    larguras = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell) or cell.value is None:
                continue
            letra = get_column_letter(cell.column)
            larguras[letra] = max(larguras.get(letra, 10), len(str(cell.value)))

    for letra, tamanho in larguras.items():
        ws.column_dimensions[letra].width = min(tamanho + 3, 45)

    ws.freeze_panes = "A2"
    wb.save(_caminho_longo(caminho))


def _formatar(caminho: Path):
    wb = load_workbook(_caminho_longo(caminho))
    header_fill = PatternFill("solid", fgColor=COR_HEADER)
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=2):
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if cell.value is not None:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", wrap_text=True)

        larguras = {}
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                letra = get_column_letter(cell.column)
                larguras[letra] = max(larguras.get(letra, 10), len(str(cell.value)))

        for letra, tamanho in larguras.items():
            ws.column_dimensions[letra].width = min(tamanho + 3, 45)

        ws.freeze_panes = "C3"

    wb.save(_caminho_longo(caminho))
